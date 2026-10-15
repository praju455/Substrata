"""Staged, idempotent tabular imports with reviewable mappings and rollback."""

import base64
import csv
import hashlib
import io
import unicodedata
import zipfile
from difflib import SequenceMatcher
from math import sqrt
from typing import Literal

from openpyxl import load_workbook
from pydantic import Field, ValidationError

from substrata.platform.budgets import reserve, settle
from substrata.platform.database import tenant_transaction
from substrata.platform.domain import (
    BomItem,
    Contract,
    Entity,
    GraphData,
    Inventory,
    Link,
    Provenance,
    PurchaseOrder,
)
from substrata.platform.models import GraphRevision, WorkRecord
from substrata.platform.providers import configured_embedder
from substrata.platform.records import add, encode, require_record
from substrata.platform.store import Conflict
from substrata.seed import canonical

CONTRACTS = {
    "entities": Entity,
    "links": Link,
    "inventory": Inventory,
    "purchase_orders": PurchaseOrder,
    "boms": BomItem,
}


class TabularImport(Contract):
    filename: str = Field(min_length=1, max_length=256)
    content_base64: str = Field(min_length=1, max_length=4_000_000)
    target: Literal["entities", "links", "inventory", "purchase_orders", "boms"]
    mapping: dict[str, str] = Field(default_factory=dict)
    units: dict[str, str] = Field(default_factory=dict)
    provenance: Provenance
    expected_version: int = Field(ge=0)
    schema_sha256: str | None = None


def normalise(name):
    return " ".join(
        "".join(
            ch if ch.isalnum() else " " for ch in unicodedata.normalize("NFKC", name).casefold()
        ).split()
    )


def parse_table(filename, content):
    if len(content) > 3_000_000:
        raise ValueError("Uploads must be at most 3 MB.")
    if filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        columns = reader.fieldnames or []
        rows = list(reader)
    elif filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 20_000_000:
                raise ValueError("Expanded workbook exceeds 20 MB.")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            columns = [str(v) if v is not None else "" for v in next(iterator, [])]
            rows = []
            for index, values in enumerate(iterator):
                if index >= 10000:
                    raise ValueError("Import contains more than 10000 rows.")
                rows.append(dict(zip(columns, values, strict=False)))
        finally:
            workbook.close()
    else:
        raise ValueError("Use CSV or XLSX; legacy binary Excel files are not supported.")
    if not columns or any(not c.strip() for c in columns) or len(set(columns)) != len(columns):
        raise ValueError("Columns must be nonblank and unique.")
    if len(rows) > 10000 or len(columns) > 100:
        raise ValueError("Import exceeds 10000 rows or 100 columns.")
    if any(None in row for row in rows):
        raise ValueError("CSV row contains more values than its header.")
    return columns, rows


def describe(body: TabularImport):
    try:
        content = base64.b64decode(body.content_base64, validate=True)
    except ValueError:
        raise ValueError("Upload is not valid base64.") from None
    columns, rows = parse_table(body.filename, content)
    schema = hashlib.sha256(canonical(columns).encode()).hexdigest()
    return content, columns, rows, schema


def validate_mapping(mapping):
    for field in mapping:
        parts = field.split(".")
        if len(parts) > 3 or any(not part or part.startswith("_") for part in parts):
            raise ValueError("Mapping path is invalid.")
        if any(".".join(parts[:length]) in mapping for length in range(1, len(parts))):
            raise ValueError("Mapping paths must not overlap a field and its nested values.")


def mapped_record(row, body):
    validate_mapping(body.mapping)
    values = (
        {"provenance": body.provenance.model_dump(mode="json")}
        if body.target in {"entities", "links"}
        else {}
    )
    for field, column in body.mapping.items():
        if column not in row:
            raise ValueError(f"Mapped column is missing: {column}")
        value = row[column]
        if value is None or value == "":
            continue
        parts = field.split(".")
        target = values
        for part in parts[:-1]:
            target = target.setdefault(part, {})
            if not isinstance(target, dict):
                raise ValueError("Mapping path overlaps a scalar value.")
        if field in body.units:
            value = {
                "value": float(value),
                "unit": body.units[field],
                "provenance": body.provenance.model_dump(mode="json"),
            }
        target[parts[-1]] = value
    return CONTRACTS[body.target].model_validate(values)


def record_key(target, row):
    if target in {"entities", "links", "purchase_orders"}:
        return row["id"]
    if target == "inventory":
        return (row["material_id"], row["plant_id"])
    return (row["material_id"], row["product_id"])


def _embedding_similarity(left, right):
    denominator = sqrt(sum(value * value for value in left) * sum(value * value for value in right))
    return (
        sum(a * b for a, b in zip(left, right, strict=True)) / denominator if denominator else 0.0
    )


def merge_entity(existing, incoming):
    # Optional dictionaries default to empty when their columns were not mapped.
    # Merge individual supplied facts so reviewing an alias cannot erase known
    # measurements (and their provenance), identifiers or attributes.
    return {
        **existing,
        **incoming,
        "id": existing["id"],
        **{
            field: {**existing[field], **incoming[field]}
            for field in ("identifiers", "attributes", "measures")
        },
    }


def stage(store, tenant_id, body: TabularImport, actor, *, embedder=None):
    content, columns, rows, schema = describe(body)
    if not body.mapping:
        return {"columns": columns, "sample": rows[:5], "schema_sha256": schema, "rows": len(rows)}
    validate_mapping(body.mapping)
    if body.schema_sha256 is not None and schema != body.schema_sha256:
        raise Conflict("Schema drift detected; inspect columns and confirm a new mapping.")
    fingerprint = hashlib.sha256(
        canonical(
            {
                "content_sha256": hashlib.sha256(content).hexdigest(),
                "target": body.target,
                "mapping": body.mapping,
                "units": body.units,
                "provenance": body.provenance.model_dump(mode="json"),
            }
        ).encode()
    ).hexdigest()
    with tenant_transaction(store.engine, tenant_id) as session:
        org = store.organisation(session, tenant_id)
        previous = session.get(WorkRecord, (tenant_id, f"import-{fingerprint}"))
        if previous:
            return encode(previous)
        if org.version != body.expected_version:
            raise Conflict("Graph changed; restart import preview.")
        base = session.get(GraphRevision, (tenant_id, org.version))
        data = base.payload if base else GraphData().model_dump(mode="json")
        version = org.version

    # Do provider work outside the graph transaction. Recheck the version before
    # writing so embeddings never hold a tenant lock or stage against stale data.
    parsed, errors, reviews = {}, [], []
    for index, row in enumerate(rows, start=2):
        if not any(value not in {None, ""} for value in row.values()):
            continue
        try:
            item = mapped_record(row, body).model_dump(mode="json")
            key = record_key(body.target, item)
            if key in parsed and parsed[key] != item:
                raise ValueError("Conflicting duplicate identity in upload.")
            parsed[key] = item
        except (ValueError, ValidationError, TypeError) as exc:
            errors.append({"row": index, "reason": str(exc)[:1000]})

    incoming_entities = []
    if body.target == "entities":
        known = {value["id"]: Entity.model_validate(value) for value in data["entities"]}
        for id in known.keys() & parsed.keys():
            parsed[id] = merge_entity(known[id].model_dump(mode="json"), parsed[id])
            known[id] = Entity.model_validate(parsed[id])
        # Only offer earlier new IDs as in-upload candidates. This makes the
        # review graph acyclic and the survivor independent of CSV row order.
        for id in sorted(parsed.keys() - known.keys()):
            item = parsed[id]
            if known:
                lexical = [
                    (e, SequenceMatcher(None, normalise(e.name), normalise(item["name"])).ratio())
                    for e in known.values()
                    if e.kind == item["kind"]
                ]
                shortlisted = sorted(
                    ((entity, score) for entity, score in lexical if score >= 0.35),
                    key=lambda pair: (-pair[1], pair[0].id),
                )[:100]
                if shortlisted:
                    incoming_entities.append((item, shortlisted))
            known[id] = Entity.model_validate(item)

    if incoming_entities:
        embedder = embedder or configured_embedder()
        pair_count = sum(len(candidates) for _, candidates in incoming_entities)
        if pair_count > 20000:
            raise ValueError("Entity resolution exceeds 20000 candidate pairs; split the import.")
        names = list(
            dict.fromkeys(
                [item["name"] for item, _ in incoming_entities]
                + [entity.name for _, candidates in incoming_entities for entity, _ in candidates]
            )
        )
        reservation = None
        if hasattr(embedder, "maximum_cost"):
            if not hasattr(actor, "tenant_id") or actor.tenant_id != tenant_id:
                raise ValueError("Live entity embeddings require authenticated tenant context.")
            reservation = reserve(
                store, actor, embedder.maximum_cost(names), "import.entity_resolution"
            )
        try:
            vectors = embedder.embed(names)
            if len(vectors) != len(names):
                raise ValueError("Embedding provider returned the wrong vector count.")
            vector_by_name = dict(zip(names, vectors, strict=True))
            for item, candidates in incoming_entities:
                ranked = []
                for entity, lexical_score in candidates:
                    semantic_score = _embedding_similarity(
                        vector_by_name[item["name"]], vector_by_name[entity.name]
                    )
                    # Keep the string-rule baseline authoritative unless semantic
                    # evidence independently crosses a conservative review threshold.
                    score = max(lexical_score, semantic_score)
                    if lexical_score >= 0.8 or semantic_score >= 0.84:
                        ranked.append(
                            {
                                "id": entity.id,
                                "name": entity.name,
                                "score": round(score, 4),
                                "name_rule_score": round(lexical_score, 4),
                                "embedding_score": round(semantic_score, 4),
                            }
                        )
                ranked.sort(key=lambda match: (-match["score"], match["id"]))
                if ranked:
                    reviews.append({"incoming_id": item["id"], "candidates": ranked[:5]})
        finally:
            if reservation:
                settle(
                    store,
                    actor,
                    reservation,
                    getattr(embedder, "telemetry", {}).get("estimated_cost_usd"),
                )

    merged = {record_key(body.target, r): r for r in data[body.target]}
    merged.update(parsed)
    candidate = {**data, body.target: list(merged.values())}
    try:
        GraphData.model_validate(candidate)
    except ValidationError as exc:
        errors.append({"row": None, "reason": str(exc)[:2000]})
    status = "invalid" if errors else "needs_review" if reviews else "staged"
    payload = {
        "filename": body.filename,
        "target": body.target,
        "columns": columns,
        "schema_sha256": schema,
        "content_sha256": hashlib.sha256(content).hexdigest(),
        "mapping": body.mapping,
        "expected_version": body.expected_version,
        "row_count": len(rows),
        "deduped_count": len(parsed),
        "errors": errors,
        "reviews": reviews,
        "candidate": candidate,
        "before_version": version,
    }
    with tenant_transaction(store.engine, tenant_id) as session:
        org = store.organisation(session, tenant_id)
        previous = session.get(WorkRecord, (tenant_id, f"import-{fingerprint}"))
        if previous:
            return encode(previous)
        if org.version != version:
            raise Conflict("Graph changed during entity resolution; restart import preview.")
        record = add(session, tenant_id, "import", status, payload, f"import-{fingerprint}")
        store.append(
            session,
            org,
            "import.staged",
            actor.subject if hasattr(actor, "subject") else actor,
            {"id": record.id, "status": status, "source_sha256": payload["content_sha256"]},
        )
        return encode(record)


def resolve_import(store, tenant_id, id, decisions: dict[str, str | None], actor, reason):
    if not reason.strip():
        raise ValueError("An entity-resolution reason is required.")
    with tenant_transaction(store.engine, tenant_id) as session:
        org = store.organisation(session, tenant_id)
        record = require_record(session, tenant_id, id, "import")
        if record.status != "needs_review":
            raise Conflict("Import is not awaiting entity review.")
        payload = dict(record.payload)
        ids = {r["incoming_id"] for r in payload["reviews"]}
        if set(decisions) != ids:
            raise ValueError("Resolve every ambiguous identity before applying the import.")
        candidate = payload["candidate"]
        entities = {e["id"]: e for e in candidate["entities"]}
        for review in payload["reviews"]:
            incoming = review["incoming_id"]
            target = decisions[incoming]
            if target is None:
                continue
            if target not in {m["id"] for m in review["candidates"]}:
                raise ValueError(
                    "Resolution must select a reviewed candidate or create a new entity."
                )

        survivors = {}
        for incoming in sorted(decisions):
            target = incoming
            visited = set()
            while decisions.get(target) is not None:
                if target in visited:
                    raise ValueError("Entity-resolution decisions must not form a cycle.")
                visited.add(target)
                target = decisions[target]
            survivors[incoming] = target
        # Resolve all paths before deleting any aliases: a reviewed candidate
        # may itself be merged into an earlier survivor in this upload.
        for incoming, target in survivors.items():
            if incoming == target:
                continue
            entities[target] = merge_entity(entities[target], entities[incoming])
            del entities[incoming]
        candidate = {**candidate, "entities": list(entities.values())}
        GraphData.model_validate(candidate)
        record.payload = {
            **payload,
            "candidate": candidate,
            "resolution": decisions,
            "resolution_reason": reason,
        }
        record.status = "staged"
        store.append(
            session, org, "import.identities.resolved", actor, {"id": id, "reason": reason}
        )
        return encode(record)


def apply_import(store, tenant_id, id, actor):
    with tenant_transaction(store.engine, tenant_id) as session:
        org = store.organisation(session, tenant_id)
        record = require_record(session, tenant_id, id, "import")
        if record.status == "applied":
            return encode(record)
        if record.status != "staged":
            raise Conflict("Only validated, reviewed imports can be applied.")
        snapshot = store.save_in(
            session,
            org,
            GraphData.model_validate(record.payload["candidate"]),
            actor,
            expected_version=record.payload["expected_version"],
            event_type="import.applied",
        )
        record.payload = {**record.payload, "applied_version": snapshot.version}
        record.status = "applied"
        return encode(record)


def rollback(store, tenant_id, id, actor, reason):
    if not reason.strip():
        raise ValueError("Rollback requires a reason.")
    with tenant_transaction(store.engine, tenant_id) as session:
        org = store.organisation(session, tenant_id)
        record = require_record(session, tenant_id, id, "import")
        if record.status != "applied" or org.version != record.payload["applied_version"]:
            raise Conflict(
                "Rollback would overwrite later changes; restore a reviewed version instead."
            )
        before = session.get(GraphRevision, (tenant_id, record.payload["before_version"]))
        data = GraphData.model_validate(before.payload) if before else GraphData()
        store.save_in(
            session, org, data, actor, expected_version=org.version, event_type="import.rolled_back"
        )
        record.status = "rolled_back"
        record.payload = {
            **record.payload,
            "rollback_reason": reason,
            "rollback_version": org.version,
        }
        return encode(record)
